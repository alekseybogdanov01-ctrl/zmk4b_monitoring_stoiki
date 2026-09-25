"""Собирает презентацию .pptx по требованиям раздела 4 ТЗ.

Содержание закрывает все обязательные пункты: концепция и архитектура,
обоснование технологий и модели, демонстрация прототипа на скриншотах, схема
связи снимков с календарным графиком, пути масштабирования и рекомендации по
установке камер.

Запуск: py -m scripts.build_presentation
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable, List, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SHOTS_DIR = PROJECT_ROOT / "docs" / "скриншоты"
TARGET = PROJECT_ROOT / "docs" / "Презентация.pptx"

BG = RGBColor(0x0B, 0x0E, 0x14)
PANEL = RGBColor(0x14, 0x1A, 0x24)
LINE = RGBColor(0x26, 0x30, 0x40)
TEXT = RGBColor(0xE8, 0xEB, 0xF0)
MUTED = RGBColor(0x96, 0xA1, 0xB2)
ACCENT = RGBColor(0x4D, 0x92, 0xF8)
DANGER = RGBColor(0xE6, 0x39, 0x46)
WARN = RGBColor(0xF4, 0xA2, 0x61)
OK = RGBColor(0x3F, 0x9E, 0x74)

FONT = "Segoe UI"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.72)


def _blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def _textbox(
    slide,
    text: str,
    *,
    left: Emu,
    top: Emu,
    width: Emu,
    height: Emu,
    size: float,
    color: RGBColor = TEXT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    line_spacing: float = 1.2,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0

    for i, line in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = FONT
    return box


def _slide_title(slide, title: str, kicker: str | None = None) -> Emu:
    top = Inches(0.55)
    if kicker:
        _textbox(
            slide,
            kicker.upper(),
            left=MARGIN,
            top=top,
            width=SLIDE_W - MARGIN * 2,
            height=Inches(0.3),
            size=11,
            color=ACCENT,
            bold=True,
        )
        top = Inches(0.92)
    _textbox(
        slide,
        title,
        left=MARGIN,
        top=top,
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.8),
        size=30,
        bold=True,
    )
    return top + Inches(1.0)


def _panel(slide, left: Emu, top: Emu, width: Emu, height: Emu, accent=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = accent or LINE
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    shape.adjustments[0] = 0.045
    shape.text_frame.text = ""
    return shape


def _bullets(slide, items: Sequence[str], *, top: Emu, size: float = 15.5) -> None:
    box = slide.shapes.add_textbox(
        MARGIN, top, SLIDE_W - MARGIN * 2, SLIDE_H - top - Inches(0.6)
    )
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0

    for i, item in enumerate(items):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.line_spacing = 1.3
        para.space_after = Pt(11)
        marker = para.add_run()
        marker.text = "— "
        marker.font.size = Pt(size)
        marker.font.color.rgb = ACCENT
        marker.font.name = FONT

        head, _, tail = item.partition(" | ")
        run = para.add_run()
        run.text = head
        run.font.size = Pt(size)
        run.font.bold = bool(tail)
        run.font.color.rgb = TEXT
        run.font.name = FONT
        if tail:
            rest = para.add_run()
            rest.text = f" {tail}"
            rest.font.size = Pt(size)
            rest.font.color.rgb = MUTED
            rest.font.name = FONT


def _cards(slide, cards: Sequence[tuple[str, str, RGBColor]], *, top: Emu) -> None:
    gap = Inches(0.3)
    total = SLIDE_W - MARGIN * 2
    width = int((total - gap * (len(cards) - 1)) / len(cards))
    height = Inches(2.5)

    for i, (title, body, accent) in enumerate(cards):
        left = MARGIN + (width + gap) * i
        _panel(slide, left, top, width, height, accent=accent)
        _textbox(
            slide,
            title,
            left=left + Inches(0.28),
            top=top + Inches(0.26),
            width=width - Inches(0.56),
            height=Inches(0.5),
            size=16,
            bold=True,
        )
        _textbox(
            slide,
            body,
            left=left + Inches(0.28),
            top=top + Inches(0.85),
            width=width - Inches(0.56),
            height=height - Inches(1.1),
            size=12.5,
            color=MUTED,
            line_spacing=1.3,
        )


def _table(slide, headers: Sequence[str], rows: Sequence[Sequence[str]], *, top: Emu):
    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(headers),
        MARGIN,
        top,
        SLIDE_W - MARGIN * 2,
        Inches(0.42) * (len(rows) + 1),
    )
    table = shape.table
    table.first_row = True

    for col, title in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = title
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x1C, 0x25, 0x33)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size = Pt(12.5)
        para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = TEXT
        para.runs[0].font.name = FONT

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(11.5)
            para.runs[0].font.color.rgb = TEXT if c == 0 else MUTED
            para.runs[0].font.bold = c == 0
            para.runs[0].font.name = FONT
    return table


def _image_slide(prs, title: str, kicker: str, image: Path, caption: str) -> None:
    slide = _blank(prs)
    _slide_title(slide, title, kicker)
    if not image.is_file():
        _textbox(
            slide,
            f"Скриншот не найден: {image.name}",
            left=MARGIN,
            top=Inches(2),
            width=SLIDE_W - MARGIN * 2,
            height=Inches(0.5),
            size=14,
            color=WARN,
        )
        return

    from PIL import Image as PILImage

    with PILImage.open(image) as img:
        ratio = img.height / img.width

    max_w = SLIDE_W - MARGIN * 2
    max_h = Inches(4.45)
    width = max_w
    height = int(width * ratio)
    if height > max_h:
        height = int(max_h)
        width = int(height / ratio)

    left = int((SLIDE_W - width) / 2)
    slide.shapes.add_picture(str(image), left, Inches(1.72), width, height)
    _textbox(
        slide,
        caption,
        left=MARGIN,
        top=Inches(1.72) + height + Inches(0.18),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.6),
        size=12.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def _arrow(slide, x1: Emu, y1: Emu, x2: Emu, y2: Emu) -> None:
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    connector.line.color.rgb = ACCENT
    connector.line.width = Pt(1.75)


# ── Слайды ─────────────────────────────────────────────────────────


def slide_title(prs) -> None:
    slide = _blank(prs)
    _textbox(
        slide,
        "BUILD WATCH",
        left=MARGIN,
        top=Inches(2.25),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.4),
        size=13,
        color=ACCENT,
        bold=True,
    )
    _textbox(
        slide,
        "Автоматизированный сервис поиска\nнарушений на строительных площадках",
        left=MARGIN,
        top=Inches(2.72),
        width=Inches(9.6),
        height=Inches(1.9),
        size=40,
        bold=True,
        line_spacing=1.12,
    )
    _textbox(
        slide,
        "Снимки с камер видеонаблюдения сверяются с календарным планом СМР "
        "автоматически: система находит технику, определяет этап работ и объясняет "
        "каждое отклонение.",
        left=MARGIN,
        top=Inches(4.75),
        width=Inches(9.2),
        height=Inches(1.1),
        size=15,
        color=MUTED,
        line_spacing=1.35,
    )
    _textbox(
        slide,
        "Техническое задание Департамента градостроительной политики города Москвы · 2026",
        left=MARGIN,
        top=Inches(6.35),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.4),
        size=12,
        color=MUTED,
    )


def slide_problem(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Что не работает сегодня", "задача")
    _cards(
        slide,
        [
            (
                "Контроль вручную",
                "Инженер открывает камеры по очереди, звонит подрядчику и оценивает "
                "ситуацию на глаз. На объект уходят часы, вывод остаётся субъективным.",
                LINE,
            ),
            (
                "Отклонения видны поздно",
                "Простой техники замечают, когда сроки уже сорваны. Управление "
                "получается реактивным: разбираются со следствиями, а не с причинами.",
                LINE,
            ),
            (
                "Данные уже есть",
                "Камеры на площадках стоят и снимают объективную картину: какая техника "
                "на объекте и когда. Этот поток просто не анализируется.",
                ACCENT,
            ),
        ],
        top=top,
    )
    _textbox(
        slide,
        "Ключевая ценность решения — доказать, что сопоставление кадров с графиком "
        "работ автоматизируется, а предупреждения об отклонениях можно объяснить.",
        left=MARGIN,
        top=top + Inches(2.85),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.8),
        size=14.5,
        color=TEXT,
        line_spacing=1.3,
    )


def slide_concept(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Концепция: три потока и один вывод", "решение")

    sources = [
        ("Кадры с камер", "JPG, PNG, WebP, BMP.\nУ каждого кадра — дата съёмки"),
        ("Календарный план СМР", "CSV или Excel: этап,\nдаты, зона работ"),
        ("Методика и справочник", "12 классов техники,\nправила «этап → техника»"),
    ]
    card_w = Inches(3.35)
    card_h = Inches(1.5)
    for i, (title, body) in enumerate(sources):
        card_top = top + (card_h + Inches(0.28)) * i
        _panel(slide, MARGIN, card_top, card_w, card_h)
        _textbox(
            slide,
            title,
            left=MARGIN + Inches(0.26),
            top=card_top + Inches(0.22),
            width=card_w - Inches(0.5),
            height=Inches(0.4),
            size=14.5,
            bold=True,
        )
        _textbox(
            slide,
            body,
            left=MARGIN + Inches(0.26),
            top=card_top + Inches(0.7),
            width=card_w - Inches(0.5),
            height=Inches(0.7),
            size=11.5,
            color=MUTED,
            line_spacing=1.25,
        )
        _arrow(
            slide,
            MARGIN + card_w,
            card_top + int(card_h / 2),
            MARGIN + card_w + Inches(0.75),
            top + Inches(2.55),
        )

    core_left = MARGIN + card_w + Inches(0.75)
    core_w = Inches(3.5)
    _panel(slide, core_left, top + Inches(1.45), core_w, Inches(2.2), accent=ACCENT)
    _textbox(
        slide,
        "Ядро анализа",
        left=core_left + Inches(0.3),
        top=top + Inches(1.68),
        width=core_w - Inches(0.6),
        height=Inches(0.4),
        size=15.5,
        bold=True,
        color=ACCENT,
    )
    _textbox(
        slide,
        "Детекция техники, агрегация\nза сутки, сопоставление\nс плановым этапом",
        left=core_left + Inches(0.3),
        top=top + Inches(2.2),
        width=core_w - Inches(0.6),
        height=Inches(1.2),
        size=12.5,
        color=MUTED,
        line_spacing=1.3,
    )
    _arrow(
        slide,
        core_left + core_w,
        top + Inches(2.55),
        core_left + core_w + Inches(0.75),
        top + Inches(2.55),
    )

    out_left = core_left + core_w + Inches(0.75)
    out_w = SLIDE_W - MARGIN - out_left
    outputs = [
        ("Статус объекта", "Простой, возможное нарушение,\nсмена стадии, норма", DANGER),
        ("Предупреждение", "Формулировка, зона работ\nи снимки-подтверждения", WARN),
        ("Таймлайн план / факт", "История по дням\nс переходом к кадру", OK),
    ]
    for i, (title, body, accent) in enumerate(outputs):
        card_top = top + (card_h + Inches(0.28)) * i
        _panel(slide, out_left, card_top, out_w, card_h, accent=accent)
        _textbox(
            slide,
            title,
            left=out_left + Inches(0.26),
            top=card_top + Inches(0.22),
            width=out_w - Inches(0.5),
            height=Inches(0.4),
            size=14.5,
            bold=True,
        )
        _textbox(
            slide,
            body,
            left=out_left + Inches(0.26),
            top=card_top + Inches(0.7),
            width=out_w - Inches(0.5),
            height=Inches(0.7),
            size=11.5,
            color=MUTED,
            line_spacing=1.25,
        )


def slide_architecture(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Архитектура и взаимодействие компонентов", "архитектура")

    blocks = [
        (
            "Frontend",
            "React 18 + Vite\nLeaflet для карты",
            "Дашборд города, карточка объекта,\nснимок с рамками, разбор площадки",
        ),
        (
            "Backend API",
            "FastAPI + Pydantic\n14 эндпоинтов",
            "Приём файлов, валидация, отдача\nсводок и таймлайнов",
        ),
        (
            "Домен",
            "Чистый Python\nбез зависимостей",
            "Методика «этап → техника»,\nотклонения, статусы",
        ),
        (
            "ML",
            "Ultralytics YOLOv8m\nOpenCV",
            "Детекция 12 классов техники,\nпороги по классу",
        ),
    ]
    gap = Inches(0.26)
    width = int((SLIDE_W - MARGIN * 2 - gap * 3) / 4)
    height = Inches(2.35)

    for i, (name, stack, purpose) in enumerate(blocks):
        left = MARGIN + (width + gap) * i
        _panel(slide, left, top, width, height, accent=ACCENT if i == 2 else LINE)
        _textbox(
            slide,
            name,
            left=left + Inches(0.24),
            top=top + Inches(0.22),
            width=width - Inches(0.48),
            height=Inches(0.4),
            size=16,
            bold=True,
        )
        _textbox(
            slide,
            stack,
            left=left + Inches(0.24),
            top=top + Inches(0.72),
            width=width - Inches(0.48),
            height=Inches(0.62),
            size=11.5,
            color=ACCENT,
            line_spacing=1.25,
        )
        _textbox(
            slide,
            purpose,
            left=left + Inches(0.24),
            top=top + Inches(1.45),
            width=width - Inches(0.48),
            height=Inches(0.8),
            size=11.5,
            color=MUTED,
            line_spacing=1.25,
        )
        if i < 3:
            _arrow(
                slide,
                left + width,
                top + int(height / 2),
                left + width + gap,
                top + int(height / 2),
            )

    _panel(slide, MARGIN, top + Inches(2.75), SLIDE_W - MARGIN * 2, Inches(1.45))
    _textbox(
        slide,
        "Хранилище: data/store.json — объекты, планы, снимки и детекции",
        left=MARGIN + Inches(0.3),
        top=top + Inches(2.98),
        width=SLIDE_W - MARGIN * 2 - Inches(0.6),
        height=Inches(0.35),
        size=14,
        bold=True,
    )
    _textbox(
        slide,
        "Отчёты за день не хранятся, а пересчитываются из плана и снимков при каждом "
        "запросе — правка методики сразу меняет выводы по всей истории. Домен не знает "
        "ни про HTTP, ни про модель: он принимает словарь «класс → количество», поэтому "
        "методика тестируется без GPU и без весов.",
        left=MARGIN + Inches(0.3),
        top=top + Inches(3.42),
        width=SLIDE_W - MARGIN * 2 - Inches(0.6),
        height=Inches(0.8),
        size=12.5,
        color=MUTED,
        line_spacing=1.3,
    )


def slide_stack(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Почему именно этот стек", "обоснование")
    _table(
        slide,
        ["Выбор", "Причина", "Альтернатива и почему отказались"],
        [
            (
                "YOLOv8m",
                "Один прогон на кадр, готовый пайплайн дообучения, работает на CPU",
                "Detectron2 точнее на мелких объектах, но тяжелее в развёртывании",
            ),
            (
                "Дообучение, не zero-shot",
                "Строительная техника в COCO отсутствует как классы",
                "Промпт-детекторы дают нестабильные имена классов",
            ),
            (
                "FastAPI",
                "Асинхронная загрузка файлов и автодокументация OpenAPI из коробки",
                "Flask потребовал бы ручных схем и валидации",
            ),
            (
                "React + Vite",
                "Нужен интерактив: карта, оверлеи рамок, таймлайн",
                "Streamlit быстрее, но не даёт нужной плотности интерфейса",
            ),
            (
                "JSON-хранилище",
                "Прототип запускается одной командой, без СУБД и миграций",
                "PostgreSQL — следующий шаг, изолирован в одном модуле",
            ),
            (
                "Порог на класс",
                "Башенный кран и грузовик распознаются хуже — им нужен мягкий порог",
                "Единый порог терял слабые классы целиком",
            ),
        ],
        top=top,
    )


def slide_pipeline(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(
        slide, "Как снимок связывается с этапом графика", "схема сопоставления"
    )

    steps = [
        ("1", "Кадр", "Дата съёмки —\nключ связи с планом"),
        ("2", "Детекция", "YOLO с порогом 0,05:\nне терять слабые объекты"),
        ("3", "Фильтр", "Свой порог уверенности\nу каждого класса"),
        ("4", "Сутки", "Детекции всех кадров\nдаты сводятся в счётчики"),
        ("5", "План на дату", "Какой этап должен идти\nи в какой зоне"),
        ("6", "Вывод", "Маркер и звено:\nстатус и отклонения"),
    ]
    gap = Inches(0.17)
    width = int((SLIDE_W - MARGIN * 2 - gap * 5) / 6)
    height = Inches(2.05)

    for i, (num, name, body) in enumerate(steps):
        left = MARGIN + (width + gap) * i
        accent = ACCENT if i in (1, 5) else LINE
        _panel(slide, left, top, width, height, accent=accent)
        _textbox(
            slide,
            num,
            left=left + Inches(0.2),
            top=top + Inches(0.18),
            width=width - Inches(0.4),
            height=Inches(0.35),
            size=12,
            color=ACCENT,
            bold=True,
        )
        _textbox(
            slide,
            name,
            left=left + Inches(0.2),
            top=top + Inches(0.58),
            width=width - Inches(0.4),
            height=Inches(0.4),
            size=14,
            bold=True,
        )
        _textbox(
            slide,
            body,
            left=left + Inches(0.2),
            top=top + Inches(1.05),
            width=width - Inches(0.4),
            height=Inches(0.9),
            size=10.5,
            color=MUTED,
            line_spacing=1.22,
        )

    example_top = top + Inches(2.5)
    _panel(slide, MARGIN, example_top, SLIDE_W - MARGIN * 2, Inches(1.75), accent=DANGER)
    _textbox(
        slide,
        "Пример из ТЗ: устройство котлована",
        left=MARGIN + Inches(0.32),
        top=example_top + Inches(0.24),
        width=SLIDE_W - MARGIN * 2 - Inches(0.64),
        height=Inches(0.4),
        size=15,
        bold=True,
    )
    _textbox(
        slide,
        "По плану 10 марта идут земляные работы в зоне «Котлован А». На кадрах за этот "
        "день система видит экскаватор — маркер этапа на месте, работы действительно "
        "начались. Но самосвалов и грузовиков в кадре нет: звено вывоза грунта разорвано. "
        "Пользователь получает предупреждение о возможном снижении темпа работ с "
        "указанием зоны и ссылками на снимки-подтверждения.",
        left=MARGIN + Inches(0.32),
        top=example_top + Inches(0.72),
        width=SLIDE_W - MARGIN * 2 - Inches(0.64),
        height=Inches(0.95),
        size=12.5,
        color=MUTED,
        line_spacing=1.3,
    )


def slide_method(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Методика «этап работ → необходимая техника»", "правила")
    _textbox(
        slide,
        "Маркер — машина, по которой этап опознаётся однозначно. Звено — набор машин, "
        "без которых работа физически не идёт: внутри группы достаточно любой машины, "
        "но все группы должны быть закрыты одновременно.",
        left=MARGIN,
        top=top - Inches(0.12),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.75),
        size=13.5,
        color=MUTED,
        line_spacing=1.3,
    )
    _table(
        slide,
        ["Этап СМР", "Маркерная техника", "Звено", "Отсутствует звено"],
        [
            (
                "Земляные работы / котлован",
                "экскаватор, бульдозер",
                "самосвал или грузовик",
                "копать можно, вывозить нечем",
            ),
            (
                "Свайный фундамент",
                "сваебой",
                "автокран или кран-манипулятор",
                "сваи нечем подавать",
            ),
            (
                "Монолитные работы",
                "бетономешалка, бетононасос",
                "бетономешалка + бетононасос",
                "насос без смеси простаивает",
            ),
            (
                "Надземная часть / монтаж",
                "башенный кран, автокран, кран-манипулятор",
                "не требуется",
                "—",
            ),
        ],
        top=top + Inches(0.85),
    )
    _textbox(
        slide,
        "Погрузчик, самосвал, грузовик и каток объявлены вспомогательными: сами по себе "
        "они не открывают этап, иначе проехавший самосвал переводил бы площадку "
        "в «земляные работы».",
        left=MARGIN,
        top=top + Inches(3.05),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.7),
        size=12.5,
        color=MUTED,
        line_spacing=1.3,
    )


def slide_deviations(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Отклонения, которые находит система", "предупреждения")
    _cards(
        slide,
        [
            (
                "Нет необходимой техники",
                "По календарю этап идёт, маркерной техники на кадрах нет. Отставание "
                "или неявка подрядной техники.\n\nКритично · missing_required",
                WARN,
            ),
            (
                "Неполное звено",
                "Маркер есть, звено разорвано: экскаватор без самосвалов, бетононасос "
                "без миксера. Формально работа идёт, фактически — простой."
                "\n\nКритично · incomplete_link",
                DANGER,
            ),
            (
                "Техника не под этап",
                "На кадре техника другого этапа. Не нарушение само по себе, но сигнал "
                "о смене стадии раньше или позже графика."
                "\n\nПредупреждение · unexpected_equipment",
                ACCENT,
            ),
        ],
        top=top,
    )
    _textbox(
        slide,
        "Каждое отклонение несёт готовую формулировку, затронутую зону из плана и "
        "ссылки на снимки — предупреждение можно проверить, а не только прочитать.",
        left=MARGIN,
        top=top + Inches(2.85),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.8),
        size=14.5,
        line_spacing=1.3,
    )


def slide_model(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Модель и качество распознавания", "машинное обучение")

    left_w = Inches(6.1)
    _table(
        slide,
        ["Параметр", "Значение"],
        [
            ("Архитектура", "YOLOv8m, дообучение из предобученных весов"),
            ("Классов техники", "12 (перечень ТЗ покрыт полностью)"),
            ("Разрешение / эпохи", "640 / 100"),
            ("Precision / Recall", "0.821 / 0.832"),
            ("mAP@50", "0.859"),
            ("mAP@50-95", "0.651"),
            ("Скорость на CPU", "медиана 0,52 с на кадр, 1,9 кадр/с"),
        ],
        top=top,
    )

    right_left = MARGIN + left_w + Inches(0.4)
    right_w = SLIDE_W - MARGIN - right_left
    _panel(slide, right_left, top, right_w, Inches(3.4), accent=ACCENT)
    _textbox(
        slide,
        "Проверочный набор собран самостоятельно",
        left=right_left + Inches(0.3),
        top=top + Inches(0.26),
        width=right_w - Inches(0.6),
        height=Inches(0.4),
        size=15,
        bold=True,
    )
    _textbox(
        slide,
        "ТЗ не даёт отдельного набора для проверки, поэтому он формируется скриптом "
        "eval_model.py: 86 кадров раскладываются по условиям съёмки и прогоняются "
        "через модель.\n\n"
        "Освещённость: яркий день, пасмурно, сумерки и контраст.\n"
        "Фон и сезон: зелёный сезон, снег и межсезонье, без зелени.\n"
        "Задействовано 10 из 12 классов техники.\n\n"
        "Скорость измеряется в рантайме: /api/detect возвращает inference_ms, "
        "/api/demo/analyze — среднее и максимум по пачке кадров.",
        left=right_left + Inches(0.3),
        top=top + Inches(0.78),
        width=right_w - Inches(0.6),
        height=Inches(2.4),
        size=12,
        color=MUTED,
        line_spacing=1.3,
    )


def slide_demo_dashboard(prs) -> None:
    _image_slide(
        prs,
        "Дашборд: что требует внимания в городе",
        "демонстрация",
        SHOTS_DIR / "дашборд.png",
        "Кадр с камеры с рамками распознанной техники, вердикт по объекту, карта "
        "и список объектов с фильтром по статусу.",
    )


def slide_demo_site(prs) -> None:
    _image_slide(
        prs,
        "Объект: план, факт и история по дням",
        "демонстрация",
        SHOTS_DIR / "объект.png",
        "Карточка объекта: пояснение отклонения, загруженные снимки, календарный план "
        "и таймлайн с переходом к конкретному кадру.",
    )


def slide_demo_method(prs) -> None:
    _image_slide(
        prs,
        "Методика доступна пользователю в интерфейсе",
        "демонстрация",
        SHOTS_DIR / "методика.png",
        "Правила сопоставления, типы отклонений и пороги по классам показаны в "
        "приложении: жюри и заказчик видят, по какой логике вынесен вердикт.",
    )


def slide_cameras(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(
        slide, "Рекомендации по установке камер на площадках", "достоверность аналитики"
    )
    _bullets(
        slide,
        [
            "Обзор всей рабочей зоны | Камера ставится так, чтобы в кадр попадал "
            "весь фронт работ, а не отдельный угол. Техника вне кадра для системы "
            "равна отсутствию техники.",
            "Высота 6–12 метров, угол наклона 15–30° | Снижает перекрытие машин друг "
            "другом и даёт силуэт целиком — именно по силуэту модель различает "
            "автокран, кран-манипулятор и башенный кран.",
            "Разрешение не ниже Full HD, техника — минимум 80 пикселей по длинной "
            "стороне | При меньшем масштабе классы начинают путаться.",
            "Одна камера на одну зону плана | Привязка отклонения к участку настолько "
            "точна, насколько зоны разделены физически.",
            "Съёмка по расписанию, 4–6 кадров за смену | Единица анализа — сутки; "
            "несколько кадров страхуют от случайного момента, когда машина уехала.",
            "Точное время в метаданных и синхронизация NTP | Дата съёмки — это ключ "
            "связи с календарным планом; сбитые часы ломают сопоставление.",
            "Стабильный ракурс и жёсткое крепление | Смена ракурса обесценивает "
            "историю наблюдений по объекту.",
            "Контровой свет, очистка объектива, подсветка | Солнце в объектив, пыль и "
            "наледь снижают полноту распознавания сильнее, чем сумерки.",
        ],
        top=top - Inches(0.1),
        size=13.5,
    )


def slide_scale(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Масштабирование и развитие", "дальше")
    _table(
        slide,
        ["Направление", "Что даёт", "Как реализуется"],
        [
            (
                "Поток вместо загрузки",
                "Контроль в режиме, близком к реальному времени",
                "Приём кадров по RTSP и очередь задач вместо ручной загрузки",
            ),
            (
                "СУБД и история",
                "Десятки тысяч объектов, аналитика по округам и подрядчикам",
                "PostgreSQL, изолирован в backend/store.py",
            ),
            (
                "Трекинг техники",
                "Учёт машино-часов, а не только факта наличия",
                "ByteTrack поверх детекций, идентификация между кадрами",
            ),
            (
                "Новые этапы и классы",
                "Отделка и инженерные системы, спецтехника заказчика",
                "Методика и справочник — это данные: добавление этапа не меняет код",
            ),
            (
                "Интеграция с ИСУП",
                "План приходит сам, отклонение уходит в задачу подрядчику",
                "Импорт графика по API вместо файла, вебхуки на отклонения",
            ),
            (
                "Дообучение на своих данных",
                "Рост качества на конкретных площадках заказчика",
                "Разметка спорных кадров из накопленной истории",
            ),
        ],
        top=top,
    )


def slide_summary(prs) -> None:
    slide = _blank(prs)
    top = _slide_title(slide, "Что закрыто по требованиям ТЗ", "итог")
    _table(
        slide,
        ["Требование раздела 2", "Реализация"],
        [
            (
                "Обнаружение и классификация техники",
                "YOLOv8m, 12 классов, mAP@50 0.859, пороги уверенности на класс",
            ),
            (
                "Сопоставление с графиком СМР",
                "Методика «этап → техника»: 4 этапа, маркеры и звенья, план из CSV и Excel",
            ),
            (
                "Визуализация результатов",
                "Дашборд города, карточка объекта, снимок с рамками, разбор площадки",
            ),
            (
                "Выявление отклонений",
                "3 типа отклонений с зоной работ и снимками-подтверждениями",
            ),
            (
                "Пример из ТЗ про котлован",
                "Экскаватор без самосвалов даёт предупреждение; закреплено тестом",
            ),
            (
                "Запуск и воспроизводимость",
                "Docker Compose, скрипт загрузки весов, тесты методики, отчёт о тестировании",
            ),
        ],
        top=top,
    )
    _textbox(
        slide,
        "Прототип запускается одной командой: docker compose up --build",
        left=MARGIN,
        top=top + Inches(3.15),
        width=SLIDE_W - MARGIN * 2,
        height=Inches(0.5),
        size=15,
        bold=True,
        color=ACCENT,
    )


SLIDES = (
    slide_title,
    slide_problem,
    slide_concept,
    slide_architecture,
    slide_stack,
    slide_pipeline,
    slide_method,
    slide_deviations,
    slide_model,
    slide_demo_dashboard,
    slide_demo_site,
    slide_demo_method,
    slide_cameras,
    slide_scale,
    slide_summary,
)


def build(target: Path) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for builder in SLIDES:
        builder(prs)
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    print(f"Презентация: {target} ({target.stat().st_size / 1024:.0f} КБ, "
          f"{len(SLIDES)} слайдов)")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=TARGET)
    args = parser.parse_args()
    build(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
