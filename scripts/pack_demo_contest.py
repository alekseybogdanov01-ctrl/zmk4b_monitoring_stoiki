"""Собрать папку «Демо_конкурс» с планом и снимками для живого показа."""

from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Демо_конкурс"

# Источники реальных кадров: сначала готовые демо-фото, потом сид рантайма.
PHOTO_SOURCES = [
    ROOT / "Демо" / "test_photos",
    ROOT / "data" / "photos",
]

# Пять кадров = пять дат с шагом 10 дней от 2025-01-10.
SHOTS = [
    {
        "file": "01_котлован_норма.jpg",
        "sites": ["jk-ok-1"],
        "date": "2025-01-10",
        "expect": "В норме",
        "why": "Экскаватор и самосвал: маркер котлована и полное звено вывоза.",
    },
    {
        "file": "02_котлован_простой.jpg",
        "sites": ["jk-idle-1"],
        "date": "2025-01-20",
        "expect": "Простой",
        "why": "Экскаватор есть, самосвала/грузовика нет — звено разорвано.",
    },
    {
        "file": "03_котлован_нет_техники.jpg",
        "sites": ["jk-warn-1"],
        "date": "2025-01-30",
        "expect": "Возможное нарушение",
        "why": "По плану котлован, маркерной техники (экскаватор/бульдозер) нет.",
    },
    {
        "file": "04_монолит_неполное_звено.jpg",
        "sites": ["jk-idle-2"],
        "date": "2025-02-09",
        "expect": "Простой",
        "why": "Бетононасос без бетономешалки — бетонное звено неполное.",
    },
    {
        "file": "05_надземка_кран.jpg",
        "sites": ["jk-ok-2", "jk-info-1"],
        "date": "2025-02-19",
        "expect": "В норме",
        "why": "Башенный кран на этапе монтажа.",
    },
]

DETECT = [
    ("экскаватор.jpg", ["jk-idle-1"]),
    ("бетононасос.jpg", ["jk-idle-2"]),
    ("башенный_кран.jpg", ["jk-ok-2", "jk-info-1"]),
]

PLAN_ROWS = [
    ("Котлован", "2025-01-10", "2025-01-31", "Котлован А"),
    ("Монолит", "2025-02-01", "2025-02-14", "Плита"),
    ("Монтаж", "2025-02-15", "2025-03-10", "Надземная часть"),
]


def find_photo(site_ids: list[str]) -> Path:
    for site_id in site_ids:
        candidates = [
            ROOT / "Демо" / "test_photos" / f"{site_id}.jpg",
            ROOT / "data" / "photos" / site_id / f"seed-{site_id}.jpg",
        ]
        photos_dir = ROOT / "data" / "photos" / site_id
        if photos_dir.is_dir():
            candidates.extend(sorted(photos_dir.glob("*.jpg")))
        train = ROOT / "data" / "train_12"
        if train.is_dir():
            candidates.extend(train.rglob(f"*{site_id}*"))
        for path in candidates:
            if path.is_file() and path.suffix.lower() in {".jpg", ".jpeg", ".png"}:
                return path
    raise FileNotFoundError(f"Нет фото для {site_ids}")


def write_plan_csv(path: Path) -> None:
    lines = ["этап,начало,конец,зона"]
    for row in PLAN_ROWS:
        lines.append(",".join(row))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def write_plan_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "plan"
    ws.append(["этап", "начало", "конец", "зона"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E8FB")
    for row in PLAN_ROWS:
        ws.append(list(row))
    for col, w in zip("ABCD", [16, 14, 14, 22]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("как_загружать")
    ws2.append(["Поле на вкладке Демо", "Значение"])
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    for row in [
        ("Календарный план", path.name),
        ("Снимки", "папка «снимки» — выделить все 5 файлов"),
        ("Базовая дата первого кадра", "2025-01-10"),
        ("Шаг между кадрами (дни)", "10"),
        ("Название объекта", "ЖК «Демо-конкурс»"),
    ]:
        ws2.append(list(row))
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 48
    for row in ws2.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_readme(path: Path) -> None:
    shots = "\n".join(
        f"  {s['file']:40} {s['date']}  → {s['expect']}\n      {s['why']}" for s in SHOTS
    )
    path.write_text(
        f"""Демо-пакет для вкладки «Демо» и «Тест модели»
===============================================

Что внутри
----------
план\\календарный_план.csv     — загрузить первым
план\\календарный_план.xlsx    — то же самое, Excel
снимки\\                       — 5 кадров таймлапса, по порядку
тест_модели\\                  — 3 отдельных кадра для вкладки «Тест модели»

Как показать анализ (вкладка «Демо»)
------------------------------------
1. Откройте http://127.0.0.1:5173/#/demo
2. «Календарный план» → календарный_план.xlsx (или .csv)
3. «Снимки с камер» → выделите все 5 файлов из папки «снимки»
   (имена начинаются с 01_, 02_… — система сортирует их по имени)
4. Базовая дата: 2025-01-10
5. Шаг между кадрами: 10
6. Название объекта: любое, например «ЖК Демо-конкурс»
7. «Запустить анализ»

Ожидаемый сюжет
---------------
{shots}

Как показать детекцию (вкладка «Тест модели»)
---------------------------------------------
Перетащите любой файл из папки «тест_модели».
На кадре появятся рамки техники.

Замечания
---------
• Это реальные кадры строительных площадок, не схемы.
• Детекцию считает текущая модель бэкенда. Если загружен fallback
  yolov8n.pt, часть техники может распознаться иначе, чем в подсказке.
• Дашборд и карта работают на заранее загруженном портфеле из 10 ЖК
  и этот пакет для них не нужен.
""",
        encoding="utf-8",
    )


def main() -> None:
    plan_dir = OUT / "план"
    shots_dir = OUT / "снимки"
    detect_dir = OUT / "тест_модели"
    for folder in (plan_dir, shots_dir, detect_dir):
        folder.mkdir(parents=True, exist_ok=True)

    write_plan_csv(plan_dir / "календарный_план.csv")
    write_plan_xlsx(plan_dir / "календарный_план.xlsx")
    write_readme(OUT / "КАК_ИСПОЛЬЗОВАТЬ.txt")

    copied = []
    for shot in SHOTS:
        src = find_photo(shot["sites"])
        dest = shots_dir / shot["file"]
        shutil.copy2(src, dest)
        copied.append(f"{dest.name} <- {src}")

    for name, sites in DETECT:
        src = find_photo(sites)
        shutil.copy2(src, detect_dir / name)

    print(f"written: {OUT}")
    for line in copied:
        print(" ", line)


if __name__ == "__main__":
    main()
