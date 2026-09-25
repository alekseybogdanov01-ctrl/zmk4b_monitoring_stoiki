"""Генерация Excel календарного плана для Демо (таймлапс каждые 10 дней).

Пять этапов: расчистка, котлован, фундаменты, каркас со стенами и перекрытиями,
благоустройство. Горизонт демо ~12 месяцев, кадр каждые 10 дней.
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Демо" / "demo_schedule_timelapse.xlsx"

# Старт строительства (первый снимок таймлапса)
START = date(2025, 1, 10)
PHOTO_STEP_DAYS = 10
# ~36 кадров ≈ 12 месяцев
PHOTO_COUNT = 36

STAGES = [
    {
        "stage": "clearing",
        "name_ru": "Расчистка участка",
        "date_from": date(2025, 1, 10),
        "date_to": date(2025, 2, 20),
        "zone": "Площадка",
        "typical_months": "1–1,5",
        "equipment": "bulldozer/loader, dump_truck/truck",
    },
    {
        "stage": "excavation",
        "name_ru": "Откопка котлована",
        "date_from": date(2025, 2, 10),
        "date_to": date(2025, 4, 20),
        "zone": "Котлован",
        "typical_months": "1,5–2,5",
        "equipment": "excavator, dump_truck/truck",
    },
    {
        "stage": "foundations",
        "name_ru": "Устройство фундаментов",
        "date_from": date(2025, 4, 1),
        "date_to": date(2025, 6, 10),
        "zone": "Фундамент",
        "typical_months": "1–2",
        "equipment": "pile_driver, autocrane/crane_manipulator",
    },
    {
        "stage": "frame",
        "name_ru": "Монтаж каркаса, стены и перекрытия",
        "date_from": date(2025, 5, 20),
        "date_to": date(2025, 11, 20),
        "zone": "Каркас",
        "typical_months": "5–8",
        "equipment": "tower_crane, concrete_mixer, concrete_pump",
    },
    {
        "stage": "landscaping",
        "name_ru": "Благоустройство",
        "date_from": date(2025, 11, 1),
        "date_to": date(2025, 12, 30),
        "zone": "Территория",
        "typical_months": "1–2",
        "equipment": "roller",
    },
]


def _primary_stage(day: date) -> tuple[str, str]:
    """При пересечении сроков берём более поздний этап графика."""
    order = ["landscaping", "frame", "foundations", "excavation", "clearing"]
    labels = {s["stage"]: s["name_ru"] for s in STAGES}
    active = [
        s["stage"]
        for s in STAGES
        if s["date_from"] <= day <= s["date_to"]
    ]
    for code in order:
        if code in active:
            return code, labels[code]
    return "", "вне активных этапов"


def build_workbook() -> Workbook:
    wb = Workbook()

    # ── Лист 1: план для загрузки в Демо ──
    ws = wb.active
    ws.title = "plan"
    headers = ["stage", "date_from", "date_to"]
    ws.append(headers)
    for s in STAGES:
        ws.append(
            [
                s["stage"],
                s["date_from"].isoformat(),
                s["date_to"].isoformat(),
            ]
        )
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E8FB")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14

    # ── Лист 2: таймлапс кадров каждые 10 дней ──
    ws2 = wb.create_sheet("timelapse_photos")
    ws2.append(
        [
            "frame",
            "captured_at",
            "expected_stage",
            "expected_stage_ru",
            "hint_equipment",
            "note",
        ]
    )
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8F5E9")

    stage_equip = {s["stage"]: s["equipment"] for s in STAGES}
    for i in range(PHOTO_COUNT):
        day = START + timedelta(days=i * PHOTO_STEP_DAYS)
        code, label = _primary_stage(day)
        ws2.append(
            [
                i + 1,
                day.isoformat(),
                code,
                label,
                stage_equip.get(code, ""),
                f"Снимок #{i + 1}, шаг {PHOTO_STEP_DAYS} дн.",
            ]
        )
    for col, w in zip("ABCDEF", [8, 14, 16, 28, 48, 24]):
        ws2.column_dimensions[col].width = w

    # ── Лист 3: пояснение по нормативам ──
    ws3 = wb.create_sheet("methodology")
    ws3.append(["Параметр", "Значение"])
    ws3["A1"].font = Font(bold=True)
    ws3["B1"].font = Font(bold=True)
    rows = [
        ("Тип объекта (референс)", "Монолитный ЖК 15–20 этажей, РФ"),
        ("Горизонт демо", f"{START.isoformat()} → {(START + timedelta(days=(PHOTO_COUNT - 1) * PHOTO_STEP_DAYS)).isoformat()}"),
        ("Шаг снимков", f"{PHOTO_STEP_DAYS} дней"),
        ("Число кадров таймлапса", str(PHOTO_COUNT)),
        ("", ""),
        ("Этап", "Типичная длительность / комментарий"),
    ]
    for r in rows:
        ws3.append(list(r))
    for s in STAGES:
        ws3.append(
            [
                s["name_ru"],
                f"{s['typical_months']} мес · {s['date_from']} → {s['date_to']}",
            ]
        )
    ws3.append(["", ""])
    ws3.append(
        [
            "Как использовать в Демо",
            "Лист plan → загрузить как календарный план; "
            "снимки именовать по порядку; "
            f"базовая дата = {START.isoformat()}, шаг = {PHOTO_STEP_DAYS}",
        ]
    )
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 72
    for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    return wb


def main() -> Path:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    wb.save(OUT)
    print(f"written: {OUT}")
    return OUT


if __name__ == "__main__":
    main()
