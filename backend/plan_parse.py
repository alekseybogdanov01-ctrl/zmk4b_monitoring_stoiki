"""Парсинг календарного плана CSV / Excel."""

from __future__ import annotations

import csv
import io
import re
from typing import Any, Dict, List, Optional

from backend.domain.stages import STAGE_LABELS_RU

# Нормализация названий этапов из русскоязычных CSV
STAGE_ALIASES: Dict[str, str] = {
    "clearing": "clearing",
    "расчистка": "clearing",
    "расчистка участка": "clearing",
    "excavation": "excavation",
    "откопка": "excavation",
    "откопка котлована": "excavation",
    "котлован": "excavation",
    "разработка котлована": "excavation",
    "устройство котлована": "excavation",
    "земляные": "excavation",
    "земляные работы": "excavation",
    "earthworks": "excavation",
    "foundations": "foundations",
    "фундамент": "foundations",
    "фундаменты": "foundations",
    "устройство фундаментов": "foundations",
    "сваи": "foundations",
    "свайные": "foundations",
    "свайный фундамент": "foundations",
    "свайные работы": "foundations",
    "piling": "foundations",
    "frame": "frame",
    "каркас": "frame",
    "монтаж каркаса": "frame",
    "монтаж каркаса, стены и перекрытия": "frame",
    "стены и перекрытия": "frame",
    "возведение стен": "frame",
    "перекрытия": "frame",
    "монолит": "frame",
    "монолитные": "frame",
    "монолитные работы": "frame",
    "бетонирование": "frame",
    "monolith": "frame",
    "надземная": "frame",
    "надземная часть": "frame",
    "монтаж": "frame",
    "монтажные работы": "frame",
    "superstructure": "frame",
    "landscaping": "landscaping",
    "благоустройство": "landscaping",
}


def normalize_stage(raw: str) -> Optional[str]:
    key = re.sub(r"\s+", " ", str(raw).strip().lower())
    if key in STAGE_ALIASES:
        return STAGE_ALIASES[key]
    # частичное совпадение
    for alias, code in STAGE_ALIASES.items():
        if alias in key or key in alias:
            return code
    return None


def _pick(row: Dict[str, str], *names: str) -> str:
    lower = {k.strip().lower(): v for k, v in row.items() if k}
    for name in names:
        if name in lower and lower[name].strip():
            return lower[name].strip()
    return ""


def parse_plan_csv(text: str) -> List[Dict[str, Any]]:
    """
    Ожидаемые колонки (любой регистр):
      stage / этап / stage_code
      date_from / начало / from
      date_to / конец / to
      zone / зона (опционально)
    """
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("CSV пуст или без заголовка")

    rows: List[Dict[str, Any]] = []
    for i, row in enumerate(reader, start=2):
        stage_raw = _pick(row, "stage", "этап", "stage_code", "код", "name", "название")
        date_from = _pick(row, "date_from", "начало", "from", "start", "date_start")
        date_to = _pick(row, "date_to", "конец", "to", "end", "date_end")
        zone = _pick(row, "zone", "зона", "участок") or None

        if not stage_raw and not date_from:
            continue

        code = normalize_stage(stage_raw)
        if code is None:
            # если уже код
            if stage_raw.strip().lower() in STAGE_LABELS_RU:
                code = stage_raw.strip().lower()
            else:
                raise ValueError(
                    f"Строка {i}: неизвестный этап «{stage_raw}». "
                    f"Допустимо: {', '.join(STAGE_LABELS_RU.keys())} "
                    f"или русские названия (расчистка, котлован, фундаменты, каркас, благоустройство)."
                )
        if not date_from or not date_to:
            raise ValueError(f"Строка {i}: нужны date_from и date_to")

        rows.append(
            {
                "stage_code": code,
                "date_from": date_from[:10],
                "date_to": date_to[:10],
                "zone": zone,
            }
        )

    if not rows:
        raise ValueError("В плане нет ни одной строки этапа")
    return rows


def parse_plan_bytes(raw: bytes, filename: str) -> List[Dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xls")):
        return _parse_excel(raw)
    # utf-8-sig для Excel-CSV
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251")
    return parse_plan_csv(text)


def _parse_excel(raw: bytes) -> List[Dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "Для Excel установите openpyxl (pip install openpyxl) или загрузите CSV"
        ) from exc

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise ValueError("Excel пуст")
    keys = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
    dict_rows: List[Dict[str, str]] = []
    for values in rows_iter:
        item = {}
        for k, v in zip(keys, values):
            item[k] = "" if v is None else str(v)
        dict_rows.append(item)

    # через CSV-хелпер
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=keys)
    writer.writeheader()
    writer.writerows(dict_rows)
    return parse_plan_csv(buf.getvalue())
